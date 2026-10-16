import csv
import io
import openpyxl
from .translation_review import output_translation

FIELDS = [('original_name','原文の検査項目名'),('name','確認・修正後の検査項目名'),
          ('translation','英訳結果'),('value','検査値'),('unit','単位'),('status','照合状態'),
          ('message','確認事項'),('original_value','抽出時の検査値'),('original_unit','抽出時の単位'),
          ('reference','基準値・基準範囲'),('original_reference','抽出時の基準値・基準範囲')]


def export(rows, extension, metadata=None):
    if extension == 'pdf':
        from .pdf_export import export_pdf
        return export_pdf(rows,metadata)
    records = [[output_translation(r) if key=='translation' else r.get(key,'') for key,_ in FIELDS] for r in rows if not r.get('excluded')]
    headers = [name for _,name in FIELDS]
    if extension == 'csv':
        stream = io.StringIO(newline='')
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(records)
        return stream.getvalue().encode('utf-8-sig')
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = '翻訳結果'
    for values in [headers] + records:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.data_type = 's'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for col in 'ABCDEFGHIJK':
        sheet.column_dimensions[col].width = 28
    if metadata:
        info = workbook.create_sheet('報告書情報')
        info.append(['項目（原文）','値（原文）','確認・修正後','項目（英訳）','値（英訳）','確認状態'])
        for item in metadata:
            info.append([item['label'],item['original_value'],item['value'],item['english_label'],item['translation'],item['status']])
        for cells in info:
            for cell in cells:
                cell.data_type = 's'
        for col in 'ABCDEF':
            info.column_dimensions[col].width = 30
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
