"""Read-only source ingestion, preserving raw text and uncertain rows."""
import csv
import io
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
import openpyxl
import pdfplumber
from PIL import Image
import pytesseract
from .normalization import key
from .report_info import extract_info


class DocumentResult(tuple):
    """Keep the legacy three-value reader API while attaching provenance."""
    def __new__(cls, rows, text, warnings, metadata, extraction):
        result = super().__new__(cls,(rows,text,warnings))
        result.metadata = metadata
        result.extraction = extraction
        return result

HEADERS = {
    'name': ['検査項目名(原文)', '検査項目名', '検査項目', '項目名'],
    'value': ['検査値(原文)', '検査値', '結果'],
    'unit': ['単位(原文)', '単位'],
}


def row(name='', value='', unit='', reason='', raw='', page=1):
    return dict(original_name=name, original_value=value, original_unit=unit,
                name=name, value=value, unit=unit, reason=reason,
                raw=raw, page=page, confirmed=not bool(reason), excluded=False,
                reference='',original_reference='',reference_present=False)


def header_mapping(values):
    result = {}
    for i, value in enumerate(values):
        for field, names in HEADERS.items():
            if key(str(value or '')) in [key(n) for n in names]:
                result[field] = i
    return result if len(result) == 3 else None


def cell_text(cell):
    if cell.data_type == 'f':
        return str(cell.value), '数式セルは対象外'
    value = cell.value
    if value is None:
        return '', ''
    if isinstance(value, str):
        return value, ''
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        fmt = cell.number_format
        if fmt == 'General':
            return str(value), ''
        if re.fullmatch(r'0(?:\.0+)?', fmt):
            digits = len(fmt.split('.')[1]) if '.' in fmt else 0
            return f'{value:.{digits}f}', ''
        return str(value), f'表示形式 {fmt} の正確な再現は要確認'
    return str(value), 'セル型の表示は要確認'


def read_excel(path, sheet):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if sheet not in workbook.sheetnames or sheet == '期待結果・検証':
            raise ValueError('入力シートを選択する。')
        result, mapping = [], None
        for cells in workbook[sheet].iter_rows():
            values = [c.value for c in cells]
            found = header_mapping(values)
            if found:
                mapping = found
                continue
            if not mapping:
                continue
            name, nr = cell_text(cells[mapping['name']])
            if not name or name.startswith(('※', '注：', '注:')):
                continue
            value, vr = cell_text(cells[mapping['value']])
            unit, ur = cell_text(cells[mapping['unit']])
            reason = ' / '.join(r for r in [nr, vr, ur] if r)
            if not value or not unit:
                reason = reason or '検査値または単位が空欄'
            result.append(row(name, value, unit, reason))
        if not mapping:
            raise ValueError('検査項目名・検査値・単位の見出しが見つからない。')
        return result
    finally:
        workbook.close()


def read_csv(path):
    content = Path(path).read_bytes()
    try:
        text = content.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = content.decode('cp932')
    result, mapping = [], None
    for values in csv.reader(io.StringIO(text)):
        found = header_mapping(values)
        if found:
            mapping = found
            continue
        if not mapping or not any(values):
            continue
        def get(field):
            return values[mapping[field]] if mapping[field] < len(values) else ''
        name, value, unit = get('name'), get('value'), get('unit')
        if not name or name.startswith(('※', '注：', '注:')):
            continue
        result.append(row(name, value, unit, '' if value and unit else '検査値または単位が空欄'))
    if not mapping:
        raise ValueError('CSVの見出しが見つからない。')
    return result


def parse_words(words, page=1, ocr=False):
    """Locate explicit table headers, then assign words by measured columns.

    Never infer missing values from the terminology master or demo workbook.
    Unparsed lines remain visible in raw text; OCR rows always require review.
    """
    lines = []
    for word in sorted(words, key=lambda w: (w['top'], w['x0'])):
        center = (word['top'] + word['bottom']) / 2
        target = next((line for y, line in lines if abs(y-center) <= max(3, (word['bottom']-word['top'])*.45)), None)
        if target is None:
            lines.append((center, [word]))
        else:
            target.append(word)
    result, columns = [], None
    for _, line in sorted(lines):
        line.sort(key=lambda w:w['x0'])
        raw = ' '.join(w['text'] for w in line)
        # Join split OCR header tokens by removing spaces only for recognition.
        compact = raw.replace(' ', '')
        if '検査項目' in compact and '単位' in compact:
            candidates = {}
            # Japanese OCR can split a header into individual characters.
            for start in range(len(line)):
                joined = ''
                for end in range(start, min(start+8,len(line))):
                    joined += line[end]['text'].replace(' ', '')
                    for label, field in [('検査項目','name'),('検査項目名','name'),('結果','value'),('検査値','value'),('単位','unit'),('基準値','reference'),('基準範囲','reference'),('判定','end')]:
                        if joined == label:
                            candidates[field] = min(candidates.get(field,float('inf')),line[start]['x0'])
            for w in line:
                text = w['text'].replace(' ', '')
                if '検査項目' in text:
                    candidates['name'] = w['x0']
                elif text in ('結果','検査値'):
                    candidates['value'] = w['x0']
                elif '単位' in text:
                    candidates['unit'] = w['x0']
                elif text.startswith('基準'):
                    candidates['reference'] = w['x0']
                elif text.startswith('判定'):
                    candidates['end'] = min(candidates.get('end', float('inf')), w['x0'])
            if all(k in candidates for k in ['name','value','unit']):
                columns = candidates if candidates['name'] < candidates['value'] < candidates['unit'] else None
            continue
        if not columns:
            continue
        if compact.startswith(('注:', '注：','※','DEMO','空欄行')):
            columns = None
            continue
        buckets = {'name':[], 'value':[], 'unit':[], 'reference':[]}
        reference_words = []
        for w in line:
            x = w['x0']
            if x < columns['name']-3 or x >= columns.get('end',float('inf'))-3:
                continue
            field = 'reference' if x >= columns.get('reference',float('inf'))-3 else 'unit' if x >= columns['unit']-3 else 'value' if x >= columns['value']-3 else 'name'
            buckets[field].append(w['text'])
            if field == 'reference':
                reference_words.append(w)
        # PDF name characters may be individually positioned; no inferred text.
        name = ''.join(buckets['name'])
        value = ''.join(buckets['value'])
        unit = ''.join(buckets['unit'])
        if not name and not value and not unit:
            continue
        reason = 'OCR抽出のため原資料との確認が必要' if ocr else ''
        if not name or not value or not unit:
            reason = '列の対応または空欄を確認'
        reference = ''
        previous = None
        for word in reference_words:
            if previous and 'x1' in previous and word['x0']-previous['x1']>1.5:
                reference += ' '
            reference += word['text']
            previous = word
        if len(reference_words)>1:
            reason = reason or '基準範囲が複数の文字断片に分かれているため原資料と確認'
        record = row(name,value,unit,reason,raw,page)
        record['reference'] = record['original_reference'] = reference
        record['reference_present'] = 'reference' in columns
        result.append(record)
    return result


def ocr_words(image):
    command = os.environ.get('TESSERACT_CMD')
    if command:
        pytesseract.pytesseract.tesseract_cmd = command
    elif Path(r'C:\Program Files\Tesseract-OCR\tesseract.exe').exists():
        pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
    try:
        langs = pytesseract.get_languages(config='')
        if not {'jpn','eng'}.issubset(langs):
            raise ValueError('日本語・英語OCRデータ（jpn / eng）が必要')
        data = pytesseract.image_to_data(image, lang='jpn+eng', config='--psm 6',
                                        output_type=pytesseract.Output.DICT, timeout=90)
    except (pytesseract.TesseractNotFoundError, pytesseract.TesseractError, RuntimeError) as exc:
        raise ValueError(f'OCRを実行できない。Tesseractとjpn/engを確認: {exc}') from exc
    return [dict(text=t, x0=data['left'][i], top=data['top'][i],
                 bottom=data['top'][i]+data['height'][i], confidence=data['conf'][i])
            for i,t in enumerate(data['text']) if t.strip()]


def read_document(path, force_ocr=False):
    path = Path(path)
    rows, texts, warnings = [], [], []
    metadata, extraction = [], []
    if path.suffix.lower() != '.pdf':
        with Image.open(path) as image:
            image.load()
            try:
                words = ocr_words(image.convert('RGB'))
                texts.append(' '.join(w['text'] for w in words))
                rows.extend(parse_words(words, ocr=True))
                metadata.extend(extract_info(texts[-1],ocr=True))
                extraction.append(dict(page=1,method='ocr',status='success' if rows else 'failed',message='' if rows else '表の列を識別できない'))
            except ValueError as exc:
                warnings.append(str(exc))
                extraction.append(dict(page=1,method='ocr',status='failed',message=str(exc)))
    else:
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > 20:
                raise ValueError('PDFは20ページ以内にする。')
            for index, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ''
                if text.strip() and not force_ocr:
                    words = page.extract_words(x_tolerance=2, y_tolerance=3)
                    texts.append(text)
                    page_rows = parse_words(words, index)
                    rows.extend(page_rows)
                    metadata.extend(extract_info(text,index))
                    state = 'success' if page_rows and not page.images else 'partial' if page_rows else 'failed'
                    message = '画像を含むページのため抽出漏れを確認' if page.images else '' if page_rows else '文字情報はあるが表を識別できない'
                    extraction.append(dict(page=index,method='text',status=state,message=message))
                    if message:
                        warnings.append(f'{index}ページ: {message}')
                else:
                    try:
                        with tempfile.TemporaryDirectory() as directory:
                            prefix = str(Path(directory)/'page')
                            exe = os.environ.get('PDFTOPPM_CMD') or shutil.which('pdftoppm')
                            if not exe:
                                raise ValueError('画像PDFのレンダリングにはPopplerのpdftoppmが必要')
                            subprocess.run([exe,'-f',str(index),'-l',str(index),'-r','200','-singlefile','-png',str(path),prefix],
                                           check=True, capture_output=True, timeout=90)
                            with Image.open(prefix+'.png') as image:
                                words = ocr_words(image)
                            texts.append(' '.join(w['text'] for w in words))
                            page_rows = parse_words(words,index,ocr=True)
                            rows.extend(page_rows)
                            # Keep line breaks from OCR coordinates for heading extraction.
                            lines = {}
                            for word in words:
                                lines.setdefault(round(word['top']/10),[]).append(word['text'])
                            metadata.extend(extract_info('\n'.join(' '.join(v) for _,v in sorted(lines.items())),index,ocr=True))
                            extraction.append(dict(page=index,method='ocr',status='partial' if page_rows else 'failed',message='OCR原文の確認が必要'))
                    except (ValueError, subprocess.SubprocessError, OSError) as exc:
                        warnings.append(f'{index}ページ: {exc}')
                        extraction.append(dict(page=index,method='ocr',status='failed',message=str(exc)))
    if not rows:
        warnings.append('表の列を識別できない。原資料と抽出文字列を確認し、行を手動追加する。')
    return DocumentResult(rows, '\n\n'.join(texts), warnings, metadata, extraction)
