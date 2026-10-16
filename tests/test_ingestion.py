import csv
import hashlib
import io
from pathlib import Path
import openpyxl
import pytest
from labtranslator.ingestion import read_excel,read_csv,read_document,parse_words,cell_text
from labtranslator.terminology import translate
from labtranslator.export import export
from labtranslator.normalization import key

DEMO=Path(__file__).resolve().parents[1]/'demo_data'


@pytest.mark.parametrize('letter,column', [('A',1),('B',2),('C',3)])
def test_excel_demo(db,letter,column):
    path=DEMO/'LabTranslator_demo_ABC.xlsx'
    before=hashlib.sha256(path.read_bytes()).hexdigest()
    records=read_excel(path,'デモ'+letter)
    book=openpyxl.load_workbook(path,read_only=True)
    expected=list(book['期待結果・検証'].values)[4:14]
    book.close()
    assert len(records)==10
    for record,expect in zip(records,expected):
        assert translate(db,record['name'])[:2]==(expect[column],'照合済み')
        assert (record['value'],record['unit'])==(expect[4],expect[5])
    assert hashlib.sha256(path.read_bytes()).hexdigest()==before


@pytest.mark.parametrize('letter',list('ABC'))
def test_pdf_text_demo(db,letter):
    records,raw,warnings=read_document(DEMO/f'LabTranslator_Demo_{letter}.pdf')
    expected=read_excel(DEMO/'LabTranslator_demo_ABC.xlsx','デモ'+letter)
    assert len(records)==10 and not warnings
    for record,expect in zip(records,expected):
        assert key(record['name'])==key(expect['name'])
        assert record['value']==expect['value']
        assert record['unit']==expect['unit']
        assert translate(db,record['name'])==translate(db,expect['name'])
    assert '受付番号' in raw


@pytest.mark.parametrize('letter',list('ABC'))
def test_png_dependency_or_extraction(letter):
    records,raw,warnings=read_document(DEMO/f'LabTranslator_Demo_{letter}.png')
    # Availability is reported separately; lack of OCR never becomes a success.
    if not records:
        assert warnings
    else:
        assert all(not r['confirmed'] for r in records)


def test_numeric_and_formula(tmp_path):
    book=openpyxl.Workbook()
    sheet=book.active
    sheet.append(['検査項目名','検査値','単位'])
    sheet.append(['血小板数',25,'×10⁴/μL'])
    sheet['B2'].number_format='0.0'
    sheet.append(['総蛋白',7.2,'g/dL'])
    sheet.append(['AST','=1+2','U/L'])
    sheet.append(['ALT',19,'U/L'])
    sheet['B5'].number_format='0.00E+00'
    path=tmp_path/'numbers.xlsx'
    book.save(path)
    records=read_excel(path,sheet.title)
    assert records[0]['value']=='25.0' and records[0]['unit']=='×10⁴/μL'
    assert records[1]['value']=='7.2'
    assert not records[2]['confirmed'] and '数式' in records[2]['reason']
    assert not records[3]['confirmed']


def test_csv_and_exports(db,tmp_path):
    path=tmp_path/'input.csv'
    path.write_text('単位,検査項目名,検査値\n×10⁴/μL,血小板数,25.0\n\n',encoding='utf-8-sig')
    records=read_csv(path)
    assert len(records)==1
    r=records[0]
    r.update(translation=translate(db,r['name'])[0],status='照合済み',message='')
    csv_data=list(csv.reader(io.StringIO(export(records,'csv').decode('utf-8-sig'))))
    assert csv_data[1][3:5]==['25.0','×10⁴/μL']
    book=openpyxl.load_workbook(io.BytesIO(export(records,'xlsx')))
    assert book.active['D2'].value=='25.0'
    assert book.active['E2'].value=='×10⁴/μL'
    r['name']='=1+1'
    book=openpyxl.load_workbook(io.BytesIO(export(records,'xlsx')))
    assert book.active['B2'].data_type=='s'


def test_layout_header_positions_and_missing_columns():
    def w(text,x,y):
        return dict(text=text,x0=x,top=y,bottom=y+10)
    words=[w('受付番号:123',0,0),w('検査項目',100,20),w('結果',300,20),w('単位',400,20),w('基準範囲',500,20),
           w('未登録検査',100,40),w('12',300,40),w('mg/dL',400,40),w('10-20',500,40),
           w('欠損',100,60),w('mg/dL',400,60),w('注：説明',100,80)]
    records=parse_words(words,ocr=True)
    assert len(records)==2 and records[0]['value']=='12'
    assert records[0]['unit']=='mg/dL' and not records[0]['confirmed']
    assert records[1]['value']=='' and '列' in records[1]['reason']
    assert parse_words([w('unknown layout',0,0)])==[]
