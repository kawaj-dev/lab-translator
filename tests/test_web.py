import io
import json
import pytest
from labtranslator import create_app


@pytest.fixture
def client(tmp_path):
    app=create_app({'TESTING':True,'DATA_DIR':tmp_path,'SECRET_KEY':'test-only'})
    with app.test_client() as client:
        client.get('/')
        yield client
    app.extensions['db_engine'].dispose()


def post(client,url,data=None,**kwargs):
    with client.session_transaction() as session:
        csrf=session['csrf']
    return client.post(url,data={'csrf':csrf,**(data or {})},**kwargs)


def upload_rows(client,monkeypatch,data):
    """Exercise review/export with controlled PDF extraction, not CSV upload."""
    import csv
    from pathlib import Path
    from labtranslator.ingestion import row
    records=[row(*values) for values in list(csv.reader(io.StringIO(data)))[1:] if values]
    monkeypatch.setattr('labtranslator.read_document',lambda path:(records,data,[]))
    pdf=Path('demo_data/LabTranslator_Demo_A.pdf').read_bytes()
    return post(client,'/',{'file':(io.BytesIO(pdf),'report.pdf')})


def test_demo_and_export(client):
    response=post(client,'/',{'demo':'C','format':'xlsx'})
    assert response.status_code==302
    location=response.headers['Location']
    html=client.get(location)
    assert 'TP(Total Protein)' in html.text
    assert '25.0' in html.text and '×10⁴/μL' in html.text
    assert client.get(location+'/export/xlsx').status_code==200
    assert client.get(location+'/export/csv').status_code==200
    assert 'シートを読み直す' not in html.text
    assert post(client,location,{'action':'sheet','sheet':'デモB'}).status_code==400
    response=post(client,'/',{'demo':'B'},follow_redirects=True)
    assert 'Serum Albumin' in response.text
    assert client.get(location+'/source').status_code==200


def test_upload_correction_and_recheck(client,monkeypatch):
    response=upload_rows(client,monkeypatch,'検査項目名,検査値,単位\n未知語,25.0,×10⁴/μL\n')
    location=response.headers['Location']
    assert '未登録' in client.get(location).text
    response=post(client,location,{'action':'save','name_0':'Alb(アルブミン)','value_0':'25.0','unit_0':'×10⁴/μL','confirmed_0':'on'},follow_redirects=True)
    assert 'Alb(Albumin)' in response.text and '未知語' in response.text
    post(client,location,{'action':'add'})
    response=post(client,location,{'action':'save','name_0':'Alb(総ビリルビン)','confirmed_0':'on','name_1':'総蛋白','value_1':'7.2','unit_1':'g/dL','confirmed_1':'on'},follow_redirects=True)
    assert '要確認' in response.text and 'Total Protein' in response.text


def test_master_and_history(client):
    assert client.get('/terms').status_code==200
    assert post(client,'/terms',{'japanese_name':'検証項目','english_name':'Test Term'},follow_redirects=True).status_code==200
    assert post(client,'/terms',{'kind':'alias','alias_name':'検証別名','term_id':'1'},follow_redirects=True).status_code==200
    assert post(client,'/terms',{'kind':'alias','alias_name':'ＴＰ','term_id':'1'}).status_code==400
    assert '検証別名' in client.get('/history').text


@pytest.mark.parametrize('extension',['pdf','png'])
def test_document_preview(client,extension):
    response=post(client,'/',{'demo':'A','format':extension})
    assert response.status_code==302
    html=client.get(response.headers['Location'])
    assert html.status_code==200 and '原資料' in html.text
    if extension=='png':
        assert 'OCR' in html.text


def test_validation(client):
    assert client.post('/',data={'demo':'A'}).status_code==400
    assert post(client,'/',{'file':(io.BytesIO(b'bad'),'bad.exe')}).status_code==400
    assert client.get('/jobs/not-a-uuid').status_code==404
    assert client.get('/logo.svg').status_code==200


def test_ocr_review_and_original_preserved(client,monkeypatch):
    from labtranslator.ingestion import row
    monkeypatch.setattr('labtranslator.read_document',lambda path:([row('総ビリルピン','0.8','mg/dL','OCR確認が必要')],'総ビリルピン 0.8 mg/dL',[]))
    response=post(client,'/',{'demo':'A','format':'png'})
    location=response.headers['Location']
    assert '要確認' in client.get(location).text
    assert 'Total Bilirubin' not in client.get(location).text
    response=post(client,location,{'action':'save','name_0':'総ビリルビン','value_0':'0.8','unit_0':'mg/dL','confirmed_0':'on'},follow_redirects=True)
    assert 'Total Bilirubin' in response.text
    assert '総ビリルピン' in response.text and '総ビリルビン' in response.text
    response=post(client,location,{'action':'save','name_0':'総ビリルビン','value_0':'','unit_0':'mg/dL','confirmed_0':'on'},follow_redirects=True)
    assert 'Total Bilirubin' not in response.text and '要確認' in response.text


def test_preview_detail_output_flow(client):
    response=post(client,'/',{'demo':'C','format':'xlsx'})
    location=response.headers['Location']
    preview=client.get(location)
    assert '日本語の原資料' in preview.text and 'LABORATORY TEST REPORT' in preview.text
    assert 'id="total-count">10' in preview.text and 'id="review-count">0' in preview.text
    assert 'id="critical-count">0' in preview.text
    assert 'id="zoom-in"' in preview.text and 'id="zoom-out"' in preview.text
    assert 'name="name_0"' not in preview.text
    assert location+'/items/0' in preview.text and location+'/output' in preview.text
    assert 'Excelをダウンロード' not in preview.text
    detail=client.get(location+'/items/0')
    assert detail.status_code==200 and '抽出原文（変更しません）' in detail.text
    response=post(client,location+'/items/0',{'name':'総蛋白','value':'7.20','unit':'g/dL','confirmed':'on'})
    assert response.status_code==302 and response.headers['Location'].endswith('#item-0')
    preview=client.get(location)
    assert 'Total Protein' in preview.text and '7.20' in preview.text
    output=client.get(location+'/output')
    assert output.status_code==200 and '出力形式を選択' in output.text
    assert location+'/export/xlsx' in output.text and location+'/export/pdf' in output.text
    assert location+'/export/csv' not in output.text
    assert '英訳が空欄になる項目が' not in output.text
    import openpyxl
    book=openpyxl.load_workbook(io.BytesIO(client.get(location+'/export/xlsx').data))
    assert book.active['C2'].value=='Total Protein' and book.active['D2'].value=='7.20'
    # Editing one row must not clear the other rows' confirmation flags.
    assert 'id="review-count">0' in preview.text
    assert client.get(location+'/items/999').status_code==404


def test_preview_warning_counts_and_blank_output(client,monkeypatch):
    data='検査項目名,検査値,単位\n未知語,1,mg/dL\nAlb(総ビリルビン),0.8,mg/dL\n'
    response=upload_rows(client,monkeypatch,data)
    location=response.headers['Location']
    preview=client.get(location)
    assert 'id="total-count">2' in preview.text
    assert 'id="review-count">2' in preview.text and 'id="critical-count">1' in preview.text
    assert '英訳未確定' in preview.text and 'Total Bilirubin' not in preview.text
    output=client.get(location+'/output')
    assert '英訳が空欄になる項目が 2 件' in output.text
    rows=list(__import__('csv').reader(io.StringIO(client.get(location+'/export/csv').data.decode('utf-8-sig'))))
    assert rows[1][2]=='' and rows[2][2]==''


def test_item_correction_exclusion_and_add(client,monkeypatch):
    from labtranslator.ingestion import row
    monkeypatch.setattr('labtranslator.read_document',lambda path:([row('総ビリルピン','0.8','mg/dL','OCR確認が必要')],'総ビリルピン 0.8 mg/dL',[]))
    location=post(client,'/',{'demo':'A','format':'png'}).headers['Location']
    post(client,location+'/items/0',{'name':'総ビリルビン','value':'0.8','unit':'mg/dL','confirmed':'on'})
    detail=client.get(location+'/items/0')
    assert '総ビリルピン' in detail.text and 'value="総ビリルビン"' in detail.text
    assert 'Total Bilirubin' in client.get(location).text
    post(client,location,{'action':'recheck'})
    assert 'Total Bilirubin' in client.get(location).text
    post(client,location+'/items/0',{'excluded':'on'})
    preview=client.get(location)
    assert '除外した行' in preview.text and 'id="total-count">0' in preview.text
    assert '出力対象の検査項目がありません' in client.get(location+'/output').text
    post(client,location+'/items/0',{'confirmed':'on'})
    assert 'Total Bilirubin' in client.get(location).text
    response=post(client,location,{'action':'add'})
    assert response.headers['Location']==location+'/items/1'
    assert client.post(location+'/items/0',data={'name':'bad'}).status_code==400


@pytest.mark.parametrize('extension',['pdf','png'])
def test_document_preview_controls(client,extension):
    location=post(client,'/',{'demo':'B','format':extension}).headers['Location']
    preview=client.get(location)
    assert 'id="source-content"' in preview.text and 'id="source-viewport"' in preview.text
    assert 'preview.js' in preview.text
    assert 'application/pdf' in preview.text and '画像としてOCRで再抽出' in preview.text
    assert 'class="source-image"' not in preview.text
    assert client.get('/static/preview.js').status_code==200


@pytest.mark.parametrize('letter',list('ABC'))
def test_simple_demo_and_pdf(client,letter):
    upload=client.get('/')
    assert 'name="format"' not in upload.text
    assert 'class="site-header"' in upload.text and 'メインナビゲーション' in upload.text
    location=post(client,'/',{'demo':letter}).headers['Location']
    assert f'LabTranslator_Demo_{letter}.pdf' in client.get(location).text
    response=client.get(location+'/export/pdf')
    assert response.status_code==200 and response.mimetype=='application/pdf'
    assert response.data.startswith(b'%PDF-')
    assert client.get(location+'/export/csv').status_code==200
    from pypdf import PdfReader
    text=''.join(page.extract_text() for page in PdfReader(io.BytesIO(response.data)).pages)
    assert 'LABORATORY TEST REPORT' in text and '25.0' in text and '×10⁴/μL' in text
    assert 'CAUTION' not in text


def test_pdf_blank_and_updated_preview_data(client,monkeypatch):
    data='検査項目名,検査値,単位\n未登録語,25.0,×10⁴/μL\n'
    location=upload_rows(client,monkeypatch,data).headers['Location']

    from pypdf import PdfReader
    def pdf_text():
        return ''.join(p.extract_text() for p in PdfReader(io.BytesIO(client.get(location+'/export/pdf').data)).pages)
    text=pdf_text()
    assert 'CAUTION' in text and 'Unregistered' in text
    assert '未登録語' not in text and '25.0' in text and '×10⁴/μL' in text
    post(client,location+'/items/0',{'name':'血小板数','value':'25.0','unit':'×10⁴/μL','confirmed':'on'})
    assert 'Platelet Count' in client.get(location).text
    assert 'Platelet Count' in pdf_text() and 'CAUTION' not in pdf_text()


@pytest.mark.parametrize('extension',['xlsx','csv','png','jpg','jpeg','exe'])
def test_pdf_only_upload(client,extension):
    response=post(client,'/',{'file':(io.BytesIO(b'%PDF-not-valid'),f'report.{extension}')})
    assert response.status_code==400
    assert 'アップロード可能な形式：PDF（第1版）' in response.text


def test_real_pdf_upload_and_invalid_contents(client):
    from pathlib import Path
    response=post(client,'/',{'file':(io.BytesIO(Path('demo_data/LabTranslator_Demo_C.pdf').read_bytes()),'report.PDF')})
    assert response.status_code==302
    assert 'TP(Total Protein)' in client.get(response.headers['Location']).text
    for data in (b'not a pdf',b'%PDF-1.7\nbroken'):
        response=post(client,'/',{'file':(io.BytesIO(data),'renamed.pdf')})
        assert response.status_code==400 and '有効なPDF' in response.text


def test_upload_design(client):
    html=client.get('/').text
    for text in ('ファイルをアップロード','英訳する検査結果報告書を選択してください。',
                 'ここにファイルをドラッグ＆ドロップ','ファイルを選択','アップロード可能な形式：PDF（第1版）'):
        assert text in html
    assert 'accept=".pdf,application/pdf"' in html and 'upload.js' in html
    assert 'name="format"' not in html and 'name="demo"' in html


@pytest.mark.parametrize('letter',list('ABC'))
def test_pdf_demo_source_metadata_and_no_excel(client,monkeypatch,letter):
    from pathlib import Path
    import hashlib
    source=Path('demo_data')/f'LabTranslator_Demo_{letter}.pdf'
    digest=hashlib.sha256(source.read_bytes()).hexdigest()
    def forbidden(*args,**kwargs):
        raise AssertionError('Excel must not be used as fallback')
    monkeypatch.setattr('labtranslator.read_excel',forbidden)
    location=post(client,'/',{'demo':letter,'format':'xlsx'}).headers['Location']
    html=client.get(location).text
    assert source.name in html and 'LabTranslator_demo_ABC.xlsx' not in html
    assert 'application/pdf' in html and 'シートを読み直す' not in html and 'name="sheet"' not in html
    assert 'Report date' in html and 'Blood (synthetic)' in html and f'DEMO-{letter}-0001' in html
    assert 'id="total-count">10' in html
    assert client.get(location+'/source').data==source.read_bytes()
    assert hashlib.sha256(source.read_bytes()).hexdigest()==digest


def test_original_pdf_image_preview(client,monkeypatch):
    from pathlib import Path
    import hashlib,shutil,os
    if not (os.environ.get('PDFTOPPM_CMD') or shutil.which('pdftoppm')):
        pytest.skip('Poppler未導入')
    path=Path('demo_data/LabTranslator_Demo_A.pdf')
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    location=post(client,'/',{'demo':'A'}).headers['Location']
    assert 'class="pdf-page-image"' in client.get(location).text
    result=client.get(location+'/source-page/1')
    assert result.status_code==200 and result.mimetype=='image/png' and result.data.startswith(b'\x89PNG')
    assert client.get(location+'/source-page/2').status_code==404
    assert client.get(location+'/source-page/0').status_code==404
    assert hashlib.sha256(path.read_bytes()).hexdigest()==digest


def test_edit_metadata_and_reference(client):
    location=post(client,'/',{'demo':'A'}).headers['Location']
    detail=client.get(location+'/metadata/2')
    assert '血液（架空）' in detail.text
    response=post(client,location+'/metadata/2',{'value':'血清','confirmed':'on'},follow_redirects=True)
    assert 'Serum' in response.text
    assert '血液（架空）' in client.get(location+'/metadata/2').text
    post(client,location+'/metadata/2',{'value':'未知検体','confirmed':'on'})
    assert '要確認' in client.get(location+'/metadata/2').text
    post(client,location+'/metadata/2',{'value':'未知検体','confirmed':'on','manual_translation':'on','english_value':'Reviewed specimen'})
    assert 'Reviewed specimen' in client.get(location).text
    post(client,location+'/items/0',{'reference':'6.5～8.0','confirmed':'on'})
    assert '6.5～8.0' in client.get(location).text
    import openpyxl
    book=openpyxl.load_workbook(io.BytesIO(client.get(location+'/export/xlsx').data))
    assert book['翻訳結果']['J2'].value=='6.5～8.0'
    assert 'Reviewed specimen' in str(list(book['報告書情報'].values))
    assert client.get(location+'/metadata/99').status_code==404
    assert client.post(location+'/metadata/2',data={}).status_code==400


def test_failed_pdf_extraction_never_falls_back(client,monkeypatch):
    from labtranslator.ingestion import DocumentResult
    monkeypatch.setattr('labtranslator.read_document',lambda path:DocumentResult([], '', ['抽出失敗'],[],[dict(page=1,method='ocr',status='failed',message='Tesseract未導入')]))
    monkeypatch.setattr('labtranslator.read_excel',lambda *args:pytest.fail('Excel fallback'))
    location=post(client,'/',{'demo':'A'}).headers['Location']
    html=client.get(location).text
    assert 'id="total-count">0' in html and '抽出失敗' in html
    assert 'Tesseract未導入' in html and 'failed' in html
    assert client.get(location+'/source').mimetype=='application/pdf'


def get_saved_job(client,location):
    from pathlib import Path
    path=Path(client.application.config['DATA_DIR'])/location.rsplit('/',1)[-1]/'job.json'
    return json.loads(path.read_text(encoding='utf-8'))


def get_pdf_text(client,location):
    from pypdf import PdfReader
    response=client.get(location+'/export/pdf')
    assert response.status_code==200
    return ''.join(p.extract_text() for p in PdfReader(io.BytesIO(response.data)).pages)


def test_manual_translation_persistence_scope_and_reset(client):
    from sqlalchemy.orm import Session
    from sqlalchemy import select,func
    from labtranslator.models import LabTerm,History
    engine=client.application.extensions['db_engine']
    with Session(engine) as db:
        before_terms=[(t.id,t.english_name) for t in db.scalars(select(LabTerm))]
        before_history=db.scalar(select(func.count()).select_from(History))
    location=post(client,'/',{'demo':'A'}).headers['Location']
    assert '>Total Protein</textarea>' in client.get(location+'/items/0').text
    result=post(client,location+'/items/0',{'english':'Protein, total (reviewed)','translation_confirmed':'on','confirmed':'on'},follow_redirects=True)
    assert 'Protein, total (reviewed)' in result.text
    post(client,location,{'action':'recheck'})
    assert 'Protein, total (reviewed)' in client.get(location).text
    assert '>Protein, total (reviewed)</textarea>' in client.get(location+'/items/0').text
    assert 'Protein, total (reviewed)' in get_pdf_text(client,location)
    import openpyxl
    book=openpyxl.load_workbook(io.BytesIO(client.get(location+'/export/xlsx').data))
    assert book.active['C2'].value=='Protein, total (reviewed)'
    another=post(client,'/',{'demo':'A'}).headers['Location']
    assert 'Protein, total (reviewed)' not in client.get(another).text
    with Session(engine) as db:
        assert before_terms==[(t.id,t.english_name) for t in db.scalars(select(LabTerm))]
        assert before_history==db.scalar(select(func.count()).select_from(History))
    post(client,location+'/items/0',{'action':'restore_translation','confirmed':'on'})
    assert get_saved_job(client,location)['rows'][0]['manual_translation'] is False
    assert 'Protein, total (reviewed)' not in client.get(location).text
    assert 'Total Protein' in client.get(location).text


@pytest.mark.parametrize('field,new_value',[('name','総タンパク'),('value','7.20'),('unit','g/L'),('reference','6.5～8.0')])
def test_source_change_keeps_manual_and_demands_reconfirmation(client,field,new_value):
    location=post(client,'/',{'demo':'A'}).headers['Location']
    post(client,location+'/items/0',{'english':'CUSTOM RETAINED','translation_confirmed':'on','confirmed':'on'})
    result=post(client,location+'/items/0',{field:new_value,'english':'CUSTOM RETAINED','translation_confirmed':'on','confirmed':'on'},follow_redirects=True)
    assert '英訳の再確認が必要' in result.text and 'CUSTOM RETAINED' in result.text
    record=get_saved_job(client,location)['rows'][0]
    assert record['manual_english']=='CUSTOM RETAINED' and not record['translation_confirmed']
    assert record['translation']=='' and record['needs_reconfirmation']
    post(client,location,{'action':'recheck'})
    assert '英訳の再確認が必要' in client.get(location).text
    output=client.get(location+'/output').text
    assert '未確認の手動英訳：1件' in output and location+'/items/0' in output
    assert '変更されたため' in output
    assert 'CUSTOM RETAINED' not in get_pdf_text(client,location)
    import openpyxl
    book=openpyxl.load_workbook(io.BytesIO(client.get(location+'/export/xlsx').data))
    assert book.active['C2'].value in (None,'')
    post(client,location+'/items/0',{'english':'CUSTOM RETAINED','translation_confirmed':'on','confirmed':'on'})
    assert 'CUSTOM RETAINED' in get_pdf_text(client,location)
    assert not get_saved_job(client,location)['rows'][0]['needs_reconfirmation']


@pytest.mark.parametrize('name,expected_status',[('未知の検査','未登録'),('Alb(総ビリルビン)','要確認')])
def test_manual_unknown_keeps_match_status(client,monkeypatch,name,expected_status):
    location=upload_rows(client,monkeypatch,f'検査項目名,検査値,単位\n{name},1,mg/dL\n').headers['Location']
    post(client,location+'/items/0',{'english':'User supplied term','confirmed':'on'})
    assert 'User supplied term' not in get_pdf_text(client,location)
    post(client,location+'/items/0',{'english':'User supplied term','translation_confirmed':'on','confirmed':'on'})
    record=get_saved_job(client,location)['rows'][0]
    assert record['match_status']==expected_status and record['status']==expected_status
    assert record['translation_confirmed'] and record['translation']=='User supplied term'
    text=get_pdf_text(client,location)
    assert 'User supplied term' in text and 'User reviewed' in text
    import openpyxl
    book=openpyxl.load_workbook(io.BytesIO(client.get(location+'/export/xlsx').data))
    assert book.active['C2'].value=='User supplied term' and book.active['F2'].value==expected_status
    assert '利用者確認済み' in book.active['G2'].value


def test_master_changes_do_not_overwrite_manual_and_restore_uses_new_candidate(client):
    location=post(client,'/',{'demo':'A'}).headers['Location']
    post(client,location+'/items/0',{'english':'Personal term','translation_confirmed':'on','confirmed':'on'})
    post(client,'/terms',{'id':'1','japanese_name':'総蛋白','english_name':'Updated standard'})
    post(client,location,{'action':'recheck'})
    assert 'Personal term' in client.get(location).text
    assert 'Updated standard' in client.get(location+'/items/0').text
    post(client,location+'/items/0',{'action':'restore_translation','confirmed':'on'})
    assert 'Updated standard' in client.get(location).text


def test_empty_manual_exclusion_and_reextraction_guard(client,monkeypatch):
    location=post(client,'/',{'demo':'A'}).headers['Location']
    post(client,location+'/items/0',{'english':'','confirmed':'on','translation_confirmed':'on'})
    record=get_saved_job(client,location)['rows'][0]
    assert record['manual_translation'] and record['translation']==''
    assert not record['translation_confirmed']
    assert '未確認の手動英訳：1件' in client.get(location+'/output').text
    post(client,location+'/items/0',{'excluded':'on','confirmed':'on'})
    assert '未確認の手動英訳：' not in client.get(location+'/output').text
    assert post(client,location,{'action':'ocr'}).status_code==400
    assert post(client,location+'/items/0',{'english':'x'*1001}).status_code==400
    assert client.post(location+'/items/0',data={'english':'No CSRF'}).status_code==400


def test_compact_preview_hides_normal_match_and_keeps_warning(client):
    location=post(client,'/',{'demo':'A'}).headers['Location']
    html=client.get(location).text
    assert '>照合済み</a>' not in html
    assert 'Reference range' in html
    detail=client.get(location+'/items/0').text
    assert '英訳（English）' in detail and 'name="translation_confirmed"' in detail
    assert 'detail-actions' in detail and 'item-detail' in detail
